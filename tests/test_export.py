import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "boss-job-filter/scripts"))
from boss_jobs import export_xlsx, report_data
from runtime import artifact_runtime
from rules import validate_config
from verify_xlsx import verify


class ExportTests(unittest.TestCase):
    def export_and_verify(self, payload):
        report = report_data(payload)
        with tempfile.TemporaryDirectory(prefix="boss-export-test-") as tmp:
            data = Path(tmp) / "data.json"
            xlsx = Path(tmp) / "test.xlsx"
            data.write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
            # Model a Codex host with no proprietary runtime installed.
            with patch('boss_jobs.artifact_runtime', return_value=None):
                export_xlsx(data, xlsx)
            return verify(xlsx, report)

    def test_empty_search_still_exports_four_sheets(self):
        payload = {"config": validate_config({"city": "深圳", "keywords": ["测试数据"]}), "collection": {"status": "result_end", "detail_status": "complete", "issues": []}, "jobs": []}
        result = self.export_and_verify(payload)
        self.assertEqual(result["total_jobs"], 0)
        self.assertEqual(result["links"], 0)

    def test_no_qualified_results_keeps_pending_excluded_and_literal_text(self):
        payload = {"config": validate_config({"mode": "市场调研", "city": "深圳", "keywords": ["测试数据"], "degree": ["本科"]}), "collection": {"status": "capture_timeout", "detail_status": "partial", "issues": ["模拟中断，仅用于测试"]}, "jobs": []}
        for n, degree in enumerate(["本科", "大专"]):
            payload["jobs"].append({"job_id": f"test{n}", "title": "=1+1", "company": "测试公司", "city": "深圳", "salary": "10-15K·13薪", "degree": degree, "experience": "", "scale": "", "jd": "", "detail_status": "failed", "detail_error": "模拟未读取", "url": f"https://www.zhipin.com/job_detail/test{n}.html", "collected_at": "2026-09-15T12:00:00+08:00"})
        result = self.export_and_verify(payload)
        self.assertEqual(result["counts"], {"符合": 0, "待确认": 1, "排除": 1})
        self.assertEqual(result["links"], 3)

    @unittest.skipUnless(artifact_runtime(), '当前设备没有 Codex 表格运行库')
    def test_codex_runtime_without_repo_node_modules(self):
        payload = {"config": validate_config({"city": "深圳", "keywords": ["测试数据"]}), "collection": {"status": "result_end", "detail_status": "complete", "issues": []}, "jobs": []}
        report = report_data(payload)
        with tempfile.TemporaryDirectory(prefix='boss 路径空格 ') as tmp:
            data, xlsx = Path(tmp) / '数据.json', Path(tmp) / '结果.xlsx'
            data.write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
            export_xlsx(data, xlsx)
            self.assertTrue(verify(xlsx, report)['verified'])

    def test_portable_full_review_preserves_jd_and_rejects_formula_injection(self):
        from rules import content_hash, digest
        from openpyxl import load_workbook
        cfg = validate_config({'city': '深圳', 'keywords': ['AI运营'], 'age': 21})
        jobs = []
        for n, salary in enumerate(['10-15K', '300元/天']):
            jobs.append({'job_id': f'review{n}', 'title': '=1+1', 'company': '测试公司', 'city': '深圳', 'salary': salary, 'degree': '本科', 'experience': '经验不限', 'scale': '20-99人', 'jd': '负责 AI 产品用户运营。年龄不限。' + '完整原文。' * 200, 'detail_status': 'complete', 'detail_fields': {}, 'url': f'https://www.zhipin.com/job_detail/review{n}.html'})
        reviews = {'config_hash': digest(cfg), 'jobs': [{'job_id': j['job_id'], 'content_hash': content_hash(j), 'checked_requirements': True, 'role': {'status': '通过', 'reason': '职责符合方向', 'evidence': '负责 AI 产品用户运营。'}, 'summary': '负责 AI 产品用户运营。', 'conflicts': []} for j in jobs]}
        report = report_data({'config': cfg, 'collection': {'status': 'target_reached', 'detail_status': 'complete'}, 'jobs': jobs}, reviews)
        with tempfile.TemporaryDirectory() as tmp:
            data, xlsx = Path(tmp) / 'data.json', Path(tmp) / 'result.xlsx'
            data.write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
            with patch('boss_jobs.artifact_runtime', return_value=None):
                export_xlsx(data, xlsx)
            self.assertEqual(verify(xlsx, report)['counts'], {'符合': 2, '待确认': 0, '排除': 0})
            wb = load_workbook(xlsx)
            self.assertEqual(wb['符合岗位']['B5'].data_type, 's')
            self.assertEqual(wb['符合岗位']['B5'].value, '=1+1')
            self.assertEqual(wb['任务说明']['B13'].value, '=SUM(B10:B12)')
            self.assertEqual(load_workbook(xlsx, data_only=True)['任务说明']['B13'].value, 2)


if __name__ == "__main__":
    unittest.main()
