"""XLSX backend for hosts without Codex's artifact runtime; no Node required."""
from datetime import datetime
import json
from pathlib import Path
import re
from xml.etree import ElementTree as ET
import zipfile

HEADERS = ['序号', '岗位名称', '岗位链接', '公司名称', '薪资原文', '学历要求', '经验要求', '企业规模', '最终结论', '判定依据 / 待确认事项', '城市 / 区域', '岗位职责摘要', '年龄要求原文', '计薪单位', '薪资下限（元）', '薪资上限（元）', '年发薪月数', '招聘者活跃状态', '行业', '福利原文', '搜索关键词', '列表采集时间', '详情采集时间', '岗位标识']
FIELDS = dict(city='城市', salary='薪资', degree='学历', experience='经验', scale='企业规模', age='年龄', detail='详情', role='岗位方向')
AUDIT = HEADERS + [label + suffix for label in FIELDS.values() for suffix in ('判定', '理由', '证据')] + ['详情读取状态', '详情失败原因', '完整职位描述', '原始岗位 URL', '实习出勤原文']
WIDTHS = [7, 34, 13, 26, 19, 14, 15, 19, 12, 60, 24, 62, 48, 12, 18, 18, 14, 20, 22, 40, 20, 24, 24, 44]
STATUSES = {'target_reached': '达到本次原始样本上限', 'result_end': '平台明确返回搜索结束', 'page_limit': '达到页数上限，覆盖未完成', 'restricted': '平台要求验证，已停止', 'login_required': '需要人工登录', 'capture_timeout': '未收到后续页面，保留部分结果', 'interrupted': '人工中断', 'browser_error': '浏览器连接或页面异常', 'partial': '部分完成', 'complete': '完成', 'not_started': '尚未开始', 'running': '执行中'}


def literal(value):
    if value is None:
        return ''
    if isinstance(value, list):
        value = '；'.join(map(str, value))
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError('单元格超过 Excel 上限，不能静默截断')
        if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
            raise ValueError('原文含 Excel 不允许的控制字符；保留原文并明确处理后重导出')
    return value


def rows_for(jobs, age_enabled, audit=False):
    for n, j in enumerate(jobs, 1):
        salary = j['salary_parsed']
        row = [n, j.get('title'), '打开岗位', j.get('company'), j.get('salary'), j.get('degree'), j.get('experience'), j.get('scale'), j['status'], j['reason'], ' / '.join(filter(None, [j.get('city'), j.get('area')])), j.get('jd_summary'), j['checks']['age'].get('evidence') or ('未写明确年龄要求' if age_enabled else '未启用年龄条件'), salary.get('unit') or '未明确', salary.get('min'), salary.get('max'), salary.get('months'), j.get('active_status') or '未获取到', j.get('industry'), j.get('welfare'), j.get('source_keyword')]
        row += [datetime.fromisoformat(j[k][:19]) if j.get(k) else '' for k in ('collected_at', 'detail_collected_at')]
        row += [j['job_id']]
        if audit:
            row += [j['checks'][key].get(field) for key in FIELDS for field in ('status', 'reason', 'evidence')]
            row += [j.get('detail_status'), j.get('detail_error'), j.get('jd'), j.get('url'), (j.get('detail_fields') or {}).get('internship_schedule')]
        yield [literal(value) for value in row]


def export(data, output):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    ink, navy, pale = '203E49', '173B4A', 'F2F7F8'
    def put(ws, row, values):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, literal(value))
            # Source text is never an executable formula, including text beginning '='.
            if isinstance(value, str):
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=10, color=ink)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    def heading(ws, row, values):
        put(ws, row, values)
        for cell in ws[row][:len(values)]:
            cell.fill = PatternFill('solid', fgColor=navy)
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        ws.row_dimensions[row].height = 36
    for index, name in enumerate(('符合岗位', '待确认岗位', '原始岗位与判定')):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = 'A66B17' if index == 1 else '117C78'
        ws.freeze_panes = 'A5'
        jobs = [j for j in data['jobs'] if index == 2 or j['status'] == ('符合' if index == 0 else '待确认')]
        headers = AUDIT if index == 2 else HEADERS
        ws.merge_cells('A1:H1')
        put(ws, 1, [f'BOSS 岗位筛选 / {name}'])
        ws['A1'].font = Font(name='Arial', size=20, bold=True, color=navy)
        ws.row_dimensions[1].height = 38
        ws.merge_cells('A2:H2')
        put(ws, 2, [f"{data['config']['mode']} · {data['config']['city']} · {' / '.join(data['config']['keywords'])} ｜ 本表 {len(jobs)} 条 · 全部样本 {len(data['jobs'])} 条"])
        ws.row_dimensions[2].height = 34
        heading(ws, 4, headers)
        for rownum, (row, job) in enumerate(zip(rows_for(jobs, data['config'].get('age') is not None, index == 2), jobs), 5):
            put(ws, rownum, row)
            if not re.fullmatch(r'https://www\.zhipin\.com/job_detail/[A-Za-z0-9_~-]+\.html', job['url']):
                raise ValueError('岗位来源链接格式不正确')
            ws.cell(rownum, 3).hyperlink = job['url']
            ws.cell(rownum, 3).font = Font(name='Arial', color='146D96', underline='single', size=10)
            ws.row_dimensions[rownum].height = 50
            if rownum % 2 == 0:
                for cell in ws[rownum]:
                    cell.fill = PatternFill('solid', fgColor=pale)
            for col in (15, 16, 17):
                ws.cell(rownum, col).number_format = '#,##0.##'
            for col in (22, 23):
                ws.cell(rownum, col).number_format = 'yyyy-mm-dd hh:mm'
        if jobs:
            table = Table(displayName=('QualifiedJobs', 'PendingJobs', 'AllJobRecords')[index], ref=f'A4:{get_column_letter(len(headers))}{len(jobs)+4}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
        else:
            ws.merge_cells('A5:H6')
            put(ws, 5, ['本次没有此类岗位。请查看“任务说明”区分筛选结果与采集状态。'])
            ws.row_dimensions[5].height = 40
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col-1] if col <= len(WIDTHS) else 52
    ws = wb.create_sheet('任务说明')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'
    ws.merge_cells('A1:B1')
    put(ws, 1, ['BOSS 岗位筛选 / 任务说明'])
    ws['A1'].font = Font(name='Arial', size=20, bold=True, color=navy)
    ws.row_dimensions[1].height = 42
    heading(ws, 3, ['项目', '内容'])
    c, s, cfg = data['collection'], data['summary'], data['config']
    compact = lambda x: '不筛选' if x is None or x == [] else json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else str(x)
    metadata = [('交付状态', s['delivery_status']), ('模式', cfg['mode']), ('城市', cfg['city']), ('岗位关键词', '、'.join(cfg['keywords'])), ('原始岗位上限', cfg['limit']), ('实际去重岗位', s['total']), ('符合岗位', s['counts']['符合']), ('待确认岗位', s['counts']['待确认']), ('排除岗位', s['counts']['排除']), ('分类合计', s['total']), ('采集状态', c.get('status')), ('详情已读取 / 未完整读取', f"{s['detail_complete']} / {s['detail_incomplete']}"), ('详情采集状态', c.get('detail_status')), ('已复核 / 尚未复核 / 方向待确认', f"{s['reviewed']} / {s['review_outstanding']} / {s['role_pending']}")]
    metadata += [(label, compact(cfg.get(key))) for key, label in [('age', '求职者年龄'), ('salary', '薪资条件（元）'), ('degree', '学历档位'), ('experience', '经验档位'), ('scale', '企业规模档位')]]
    metadata += [('薪资口径', '薪资区间有交集；计薪单位分别比较，不换算日薪或年发薪月数。区间不代表企业承诺的金额。'), ('年龄与缺失', '启用年龄条件时，只看明确任职年龄要求；缺失、含糊或冲突均待确认。'), ('结论口径', '任一条件明确不符即排除；无明确不符但有未知则待确认；全部已验证通过才符合。'), ('详情与复核', '完整职责逐条复核；正文与字段冲突按待确认处理。未完整读取不计入符合。'), ('档位口径', '学历、经验、规模只接受指定档位；明确不限与缺失分开。'), ('样本边界', s['coverage']), ('页数 / 去重 / 无效', f"{c.get('pages_received', 0)} / {c.get('duplicates', 0)} / {c.get('invalid_rows', 0)}"), ('时间口径', '表内沿用采集机器当地时间；原始 JSON 保留带时区的 ISO 时间。'), ('开始时间', c.get('started_at')), ('列表结束时间', c.get('finished_at')), ('详情结束时间', c.get('details_finished_at')), ('异常记录', '；'.join(c.get('issues', [])) or '无'), ('未执行关键词', '、'.join(k for k in cfg['keywords'] if k not in {q['keyword'] for q in c.get('queries', [])}) or '无')]
    metadata += [(f'关键词采集：{q["keyword"]}', compact(q)) for q in c.get('queries', [])]
    for n, row in enumerate(metadata, 4):
        if row[0] in ('采集状态', '详情采集状态'):
            row = (row[0], STATUSES.get(row[1], row[1]))
        put(ws, n, row)
        ws.cell(n, 1).fill = PatternFill('solid', fgColor=pale)
        ws.row_dimensions[n].height = 48
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 112
    caches = {}
    for row, status in [(10, '符合'), (11, '待确认'), (12, '排除')]:
        ws[f'B{row}'] = f'=COUNTIF(\'原始岗位与判定\'!$I$5:$I${max(5, len(data["jobs"])+4)},"{status}")'
        caches[f'B{row}'] = s['counts'][status]
    ws['B13'] = '=SUM(B10:B12)'
    caches['B13'] = s['total']
    for row in range(8, 14):
        ws.cell(row, 2).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    # openpyxl writes formulas but does not calculate them. Cache these four known
    # count results; Excel still recalculates after edits. No generic evaluator.
    ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    temp = output.with_suffix('.cache.tmp')
    try:
        with zipfile.ZipFile(output) as source, zipfile.ZipFile(temp, 'w', zipfile.ZIP_DEFLATED) as dest:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet4.xml':
                    root = ET.fromstring(content)
                    for cell in root.findall('.//x:sheetData/x:row/x:c', ns):
                        if cell.get('r') in caches:
                            cell.find('x:v', ns).text = str(caches[cell.get('r')])
                    content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                dest.writestr(item, content)
        temp.replace(output)
    finally:
        temp.unlink(missing_ok=True)
